from fastapi import FastAPI, Form, UploadFile, File
from fastapi.responses import HTMLResponse, FileResponse, RedirectResponse
from pathlib import Path
from typing import List
from uuid import uuid4
from docx import Document
from docx.shared import Inches
import shutil

app = FastAPI(title="APP DE TASACIONES RD")

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
UPLOADS = DATA / "uploads"
REPORTS = DATA / "reports"

for p in (UPLOADS, REPORTS):
    p.mkdir(parents=True, exist_ok=True)

STYLE = """
<style>
*{box-sizing:border-box}
body{
    font-family:Arial,sans-serif;
    background:#f4f6f8;
    margin:0;
    color:#172033;
}
.top{
    background:#172033;
    color:white;
    padding:18px 6%;
}
.wrap{
    max-width:1100px;
    margin:30px auto;
    padding:0 20px;
}
.card{
    background:white;
    padding:28px;
    border-radius:12px;
    box-shadow:0 2px 10px rgba(0,0,0,.08);
    margin-bottom:20px;
}
.grid{
    display:grid;
    grid-template-columns:repeat(2,1fr);
    gap:18px;
}
.full{grid-column:1/-1}
label{
    display:block;
    font-weight:bold;
    margin-bottom:6px;
}
input,select,textarea{
    width:100%;
    padding:11px;
    border:1px solid #cbd2da;
    border-radius:6px;
}
textarea{min-height:90px}
button,.btn{
    display:inline-block;
    background:#172033;
    color:white;
    padding:12px 22px;
    border:0;
    border-radius:7px;
    text-decoration:none;
    cursor:pointer;
}
.note{
    background:#eef4ff;
    padding:12px;
    border-radius:7px;
    font-size:14px;
}
.ok{
    background:#e9f8ee;
    border-left:5px solid #238636;
    padding:15px;
}
.warn{
    background:#fff6df;
    border-left:5px solid #d29922;
    padding:15px;
}
h1,h2,h3{margin-top:0}
small{color:#667085}
@media(max-width:700px){
    .grid{grid-template-columns:1fr}
    .full{grid-column:auto}
}
</style>
"""


def page(body):
    return HTMLResponse(
        f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>APP DE TASACIONES RD</title>
{STYLE}
</head>
<body>
<div class="top"><b>APP DE TASACIONES RD</b></div>
{body}
</body>
</html>"""
    )


@app.get("/", response_class=HTMLResponse)
def home():
    return page("""
<div class="wrap">
<div class="card">
<h1>Nueva tasación</h1>
<p>Complete los datos del expediente y cargue los documentos disponibles.</p>

<form action="/generar" method="post" enctype="multipart/form-data">

<h2>1. Datos generales</h2>
<div class="grid">

<div>
<label>Expediente</label>
<input name="expediente" placeholder="Ej.: TAS-2026-001">
</div>

<div>
<label>Banco</label>
<select name="banco">
<option value="">Seleccione</option>
<option>Banco Popular Dominicano</option>
<option>Popular Bank, LTD</option>
<option>Otro</option>
</select>
</div>

<div>
<label>Propietario</label>
<input name="propietario">
</div>

<div>
<label>Fecha de tasación</label>
<input type="date" name="fecha">
</div>

<div class="full">
<label>Dirección</label>
<input name="direccion">
</div>

<div>
<label>Sector</label>
<input name="sector">
</div>

<div>
<label>Municipio</label>
<input name="municipio">
</div>

<div>
<label>Provincia</label>
<input name="provincia">
</div>

</div>

<br>
<h2>2. Datos registrales</h2>

<div class="grid">

<div>
<label>Matrícula</label>
<input name="matricula">
</div>

<div>
<label>Libro</label>
<input name="libro">
</div>

<div>
<label>Folio</label>
<input name="folio">
</div>

<div>
<label>Área (m²)</label>
<input name="area">
</div>

<div class="full">
<label>Designación Catastral</label>
<input
    name="designacion_catastral"
    minlength="12"
    placeholder="Mínimo 12 caracteres"
>
<small>
Campo independiente. Puede contener más de 12 caracteres según el documento.
</small>
</div>

<div class="full">
<label>Designación Posicional</label>
<input
    name="designacion_posicional"
    inputmode="numeric"
    pattern="[0-9]{12}"
    maxlength="12"
    placeholder="12 dígitos, cuando esté disponible"
>
<small>
No confundir con la Designación Catastral. Este campo se utilizará posteriormente
para el cálculo/verificación de coordenadas.
</small>
</div>

</div>

<br>
<h2>3. Descripción de las mejoras</h2>

<div class="grid">

<div class="full">
<label>Mejoras / descripción técnica</label>
<textarea name="mejoras"
placeholder="Describa espacios, materiales, terminaciones, estado físico y demás características."></textarea>
</div>

</div>

<br>
<h2>4. Documentos y fotografías</h2>

<div class="grid">

<div class="full">
<label>Plantilla Word maestra (.docx)</label>
<input type="file" name="plantilla" accept=".docx">
<small>
En esta versión de prueba se recibe la plantilla. La integración de fidelidad total
de la plantilla se continuará desarrollando.
</small>
</div>

<div class="full">
<label>Foto de fachada seleccionada por el tasador</label>
<input type="file" name="fachada" accept="image/*">
</div>

<div class="full">
<label>Fotografías del inmueble</label>
<input type="file" name="fotos" accept="image/*" multiple>
</div>

<div class="full">
<label>Títulos / documentos legales</label>
<input type="file" name="titulos" multiple>
</div>

</div>

<br>
<div class="note">
<b>Importante:</b> esta es una versión de prueba del flujo. Los datos generados
deben ser revisados por el tasador antes de utilizarse profesionalmente.
</div>

<br>
<button type="submit">Generar tasación de prueba</button>

</form>
</div>
</div>
""")


def save_upload(up: UploadFile | None, folder: Path):
    if not up or not up.filename:
        return None

    folder.mkdir(parents=True, exist_ok=True)
    safe_name = Path(up.filename).name
    path = folder / f"{uuid4().hex[:8]}_{safe_name}"

    with path.open("wb") as f:
        shutil.copyfileobj(up.file, f)

    return path


@app.post("/generar", response_class=HTMLResponse)
def generar(
    expediente: str = Form(""),
    banco: str = Form(""),
    propietario: str = Form(""),
    fecha: str = Form(""),
    direccion: str = Form(""),
    sector: str = Form(""),
    municipio: str = Form(""),
    provincia: str = Form(""),
    matricula: str = Form(""),
    libro: str = Form(""),
    folio: str = Form(""),
    area: str = Form(""),
    designacion_catastral: str = Form(""),
    designacion_posicional: str = Form(""),
    mejoras: str = Form(""),
    plantilla: UploadFile | None = File(None),
    fachada: UploadFile | None = File(None),
    fotos: List[UploadFile] = File(default=[]),
    titulos: List[UploadFile] = File(default=[]),
):

    # Validaciones registrales básicas
    dc = designacion_catastral.strip()
    dp = designacion_posicional.strip()

    if dc and len(dc) < 12:
        return page("""
        <div class="wrap">
        <div class="card">
        <div class="warn">
        <h2>Revisar Designación Catastral</h2>
        <p>La Designación Catastral debe contener al menos 12 caracteres.</p>
        </div>
        <br>
        <a class="btn" href="/">Volver al formulario</a>
        </div>
        </div>
        """)

    if dp and (len(dp) != 12 or not dp.isdigit()):
        return page("""
        <div class="wrap">
        <div class="card">
        <div class="warn">
        <h2>Revisar Designación Posicional</h2>
        <p>La Designación Posicional debe contener exactamente 12 dígitos.</p>
        </div>
        <br>
        <a class="btn" href="/">Volver al formulario</a>
        </div>
        </div>
        """)

    eid = "".join(
        c for c in expediente if c.isalnum() or c in "-_"
    )[:50] or uuid4().hex[:8]

    folder = UPLOADS / eid
    folder.mkdir(parents=True, exist_ok=True)

    template_path = save_upload(plantilla, folder)
    fachada_path = save_upload(fachada, folder)

    photo_paths = []
    for photo in fotos:
        saved = save_upload(photo, folder)
        if saved:
            photo_paths.append(saved)

    title_paths = []
    for title in titulos:
        saved = save_upload(title, folder)
        if saved:
            title_paths.append(saved)

    # Documento de prueba.
    # La fidelidad completa de la plantilla maestra se implementará
    # en una fase posterior del desarrollo.
    doc = Document()

    doc.add_heading("INFORME DE TASACIÓN - PRUEBA", 0)

    doc.add_heading("Datos generales", level=1)
    doc.add_paragraph(f"Expediente: {expediente}")
    doc.add_paragraph(f"Banco: {banco}")
    doc.add_paragraph(f"Propietario: {propietario}")
    doc.add_paragraph(f"Fecha: {fecha}")
    doc.add_paragraph(f"Dirección: {direccion}")
    doc.add_paragraph(f"Sector: {sector}")
    doc.add_paragraph(f"Municipio: {municipio}")
    doc.add_paragraph(f"Provincia: {provincia}")

    doc.add_heading("Datos registrales", level=1)
    doc.add_paragraph(f"Matrícula: {matricula}")
    doc.add_paragraph(f"Libro: {libro}")
    doc.add_paragraph(f"Folio: {folio}")
    doc.add_paragraph(f"Área: {area} m²")
    doc.add_paragraph(f"Designación Catastral: {designacion_catastral}")
    doc.add_paragraph(f"Designación Posicional: {designacion_posicional}")

    if banco == "Banco Popular Dominicano":
        doc.add_paragraph(
            "Tasación solo es válida para el Banco Popular Dominicano"
        )

    elif banco == "Popular Bank, LTD":
        doc.add_paragraph(
            "Tasación válida solo para el Popular Bank, LTD"
        )

    if mejoras.strip():
        doc.add_heading("Análisis de mejoras", level=1)
        doc.add_paragraph(mejoras)

    if fachada_path:
        doc.add_heading("Fachada seleccionada", level=1)
        try:
            doc.add_picture(str(fachada_path), width=Inches(5.8))
        except Exception:
            doc.add_paragraph(
                "La fotografía de fachada fue cargada, pero no pudo insertarse."
            )

    if photo_paths:
        doc.add_heading("Fotografías", level=1)
        for photo in photo_paths:
            try:
                doc.add_picture(str(photo), width=Inches(4.5))
            except Exception:
                doc.add_paragraph(f"Fotografía cargada: {photo.name}")

    if title_paths:
        doc.add_page_break()
        doc.add_heading("Documentos legales cargados", level=1)

        for title in title_paths:
            doc.add_paragraph(title.name)

    if template_path:
        doc.add_paragraph(
            "Plantilla maestra recibida para procesamiento: "
            + template_path.name
        )

    output_name = f"{eid}_tasacion_prueba.docx"
    output_path = REPORTS / output_name
    doc.save(output_path)

    checks = [
        ("Propietario", bool(propietario.strip())),
        ("Dirección", bool(direccion.strip())),
        ("Matrícula", bool(matricula.strip())),
        ("Libro", bool(libro.strip())),
        ("Folio", bool(folio.strip())),
        ("Área", bool(area.strip())),
        ("Designación Catastral", bool(dc)),
        ("Fachada", bool(fachada_path)),
        ("Título / documento legal", bool(title_paths)),
    ]

    qa = sum(1 for _, ok in checks if ok)
    total = len(checks)

    rows = "".join(
        f"<li>{'✅' if ok else '⚠️'} {name}</li>"
        for name, ok in checks
    )

    return page(f"""
    <div class="wrap">

    <div class="card">
    <div class="ok">
    <h2>Tasación de prueba generada</h2>
    <p>El servidor procesó correctamente el formulario.</p>
    </div>

    <br>

    <h3>Expediente</h3>
    <p><b>{expediente or eid}</b></p>

    <h3>Control preliminar</h3>
    <p><b>{qa} de {total}</b> elementos principales completados.</p>

    <ul>
    {rows}
    </ul>

    <p>
    <a class="btn" href="/descargar/{output_name}">
    Descargar Word de prueba
    </a>
    </p>

    <p>
    <a href="/">Crear otra tasación</a>
    </p>

    </div>
    </div>
    """)


# Evita mostrar "Method Not Allowed" si alguien abre /generar directamente.
@app.get("/generar")
def generar_get():
    return RedirectResponse(url="/", status_code=303)


@app.get("/descargar/{filename}")
def descargar(filename: str):
    safe = Path(filename).name
    path = REPORTS / safe

    if not path.exists():
        return page("""
        <div class="wrap">
        <div class="card">
        <div class="warn">
        <h2>Archivo no encontrado</h2>
        <p>El documento solicitado ya no está disponible.</p>
        </div>
        <br>
        <a class="btn" href="/">Volver</a>
        </div>
        </div>
        """)

    return FileResponse(
        path,
        filename=safe,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )


@app.get("/health")
def health():
    return {
        "status": "ok",
        "app": "tasaciones-rd",
        "version": "0.2-test",
    }
